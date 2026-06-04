import streamlit as st
import fitz
import os
import re
import tempfile
import io
from pdf2docx import Converter
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Các thư viện phục vụ cho PowerPoint
from pptx import Presentation
from pptx.util import Inches, Pt as PptxPt

# ── Kiểm tra PDF có footer "Trang X/Y" không ────────────────────
def pdf_has_page_marker(pdf_path, password=None):
    page_pattern = re.compile(r'Trang\s+\d+(/\d+)?', re.IGNORECASE)
    try:
        doc = fitz.open(pdf_path)
        if password:
            doc.authenticate(password)
        for page in doc:
            text = page.get_text()
            if page_pattern.search(text):
                doc.close()
                return True
        doc.close()
    except Exception:
        pass
    return False

# ── Gộp nhiều section về 1 (Cho Word) ───────────────────────────
def merge_sections_to_one(doc):
    body = doc.element.body
    removed = 0
    for para in body.findall('.//' + qn('w:p')):
        pPr = para.find(qn('w:pPr'))
        if pPr is not None:
            sectPr = pPr.find(qn('w:sectPr'))
            if sectPr is not None:
                pPr.remove(sectPr)
                removed += 1
    return removed

# ── Xóa dòng trống cuối file (Cho Word) ─────────────────────────
def remove_trailing_empty_paragraphs(doc):
    for p in reversed(doc.paragraphs):
        if not p.text.strip():
            p._element.getparent().remove(p._element)
        else:
            break

# ── Tạo field PAGE / NUMPAGES (Cho Word) ────────────────────────
def add_complex_field(run, instruction):
    r = run._r
    fc_begin = OxmlElement('w:fldChar')
    fc_begin.set(qn('w:fldCharType'), 'begin')
    r.append(fc_begin)

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    r.append(instr)

    fc_sep = OxmlElement('w:fldChar')
    fc_sep.set(qn('w:fldCharType'), 'separate')
    r.append(fc_sep)

    wt = OxmlElement('w:t')
    wt.text = '1'
    r.append(wt)

    fc_end = OxmlElement('w:fldChar')
    fc_end.set(qn('w:fldCharType'), 'end')
    r.append(fc_end)

# ── Thiết lập footer (Cho Word) ──────────────────────────────────
def set_footer(doc):
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer

        for para in footer.paragraphs:
            para._element.getparent().remove(para._element)

        fp = footer.add_paragraph()
        fp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        fp.add_run("Trang ")
        run_page = fp.add_run()
        add_complex_field(run_page, " PAGE ")
        fp.add_run("/")
        run_total = fp.add_run()
        add_complex_field(run_total, " NUMPAGES ")

        for run in fp.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.italic = True
            rPr = run._r.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:ascii"),    "Times New Roman")
            rFonts.set(qn("w:hAnsi"),    "Times New Roman")
            rFonts.set(qn("w:eastAsia"), "Times New Roman")

# ── Xóa toàn bộ footer (Cho Word) ───────────────────────────────
def clear_all_footers(doc):
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer
        for para in footer.paragraphs:
            para._element.getparent().remove(para._element)

# ── Hậu xử lý Word ──────────────────────────────────────────────
def post_process_word(docx_path, has_marker):
    doc = Document(docx_path)
    page_pattern = re.compile(r'^Trang\s+\d+(/\d+)?$', re.IGNORECASE)
    to_remove = [p for p in doc.paragraphs if page_pattern.match(p.text.strip())]
    for para in to_remove:
        para._element.getparent().remove(para._element)

    remove_trailing_empty_paragraphs(doc)
    merge_sections_to_one(doc)

    if has_marker:
        set_footer(doc)
    else:
        clear_all_footers(doc)

    doc.save(docx_path)

# ── HÀM CHUYỂN PDF SANG PPTX NÂNG CAO: GIỮ BẢNG + TEXT SỬA ĐƯỢC ──
def convert_pdf_to_pptx(pdf_path, start_page, end_page, password=None):
    prs = Presentation()
    blank_layout = prs.slide_layouts[6] # Slide trống hoàn toàn
    
    doc = fitz.open(pdf_path)
    if password:
        doc.authenticate(password)
        
    total_p = len(doc)
    p_start = max(0, start_page - 1) if start_page > 0 else 0
    p_end = min(total_p, end_page) if end_page > 0 else total_p
    
    for page_idx in range(p_start, p_end):
        pdf_page = doc[page_idx]
        page_width = pdf_page.rect.width
        page_height = pdf_page.rect.height
        
        # Thiết lập kích thước Slide khít với PDF gốc
        prs.slide_width = Inches(page_width / 72)
        prs.slide_height = Inches(page_height / 72)
        
        slide = prs.slides.add_slide(blank_layout)
        
        # 1. PHÁT HIỆN VÀ DỰNG BẢNG (TABLES) NATIVE TRÊN POWERPOINT
        tables = pdf_page.find_tables()
        table_bboxes = []
        
        for t in tables:
            data = t.extract() # Lấy mảng dữ liệu 2 chiều (rows x cols)
            if not data:
                continue
            rows = len(data)
            cols = len(data[0]) if rows > 0 else 0
            if rows == 0 or cols == 0:
                continue
            
            # Lưu lại tọa độ của bảng để lát nữa không trích xuất text đè lên vùng này
            table_bboxes.append(t.bbox)
            
            # Quy đổi tọa độ bảng sang PowerPoint
            left = Inches(t.bbox[0] / 72)
            top = Inches(t.bbox[1] / 72)
            width = Inches((t.bbox[2] - t.bbox[0]) / 72)
            height = Inches((t.bbox[3] - t.bbox[1]) / 72)
            
            # Khởi tạo bảng PowerPoint nguyên bản
            table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
            pptx_table = table_shape.table
            
            # Đổ dữ liệu chữ vào từng Cell trong bảng
            for r_idx, row_data in enumerate(data):
                for c_idx, cell_value in enumerate(row_data):
                    cell = pptx_table.cell(r_idx, c_idx)
                    cell.text = str(cell_value) if cell_value is not None else ""
                    # Định dạng font chữ trong bảng
                    for p in cell.text_frame.paragraphs:
                        p.font.name = "Times New Roman"
                        p.font.size = PptxPt(10)
                        
        # 2. TRÍCH XUẤT HÌNH ẢNH / LOGO
        image_infos = pdf_page.get_image_info(hashes=False, xrefs=True)
        for img_info in image_infos:
            xref = img_info['xref']
            if xref > 0:
                try:
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image["image"]
                    bbox = img_info['bbox']
                    
                    left = Inches(bbox[0] / 72)
                    top = Inches(bbox[1] / 72)
                    width = Inches((bbox[2] - bbox[0]) / 72)
                    height = Inches((bbox[3] - bbox[1]) / 72)
                    
                    slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width, height)
                except Exception:
                    pass
                    
        # 3. TRÍCH XUẤT VĂN BẢN (TỰ ĐỘNG BỎ QUA CÁC TEXT NẰM TRONG BẢNG)
        page_dict = pdf_page.get_text("dict")
        for block in page_dict.get("blocks", []):
            if "lines" in block:
                block_bbox = block["bbox"]
                
                # Tính toán điểm trung tâm của block text để kiểm tra xem nó có nằm lọt vào trong bảng không
                cx = (block_bbox[0] + block_bbox[2]) / 2
                cy = (block_bbox[1] + block_bbox[3]) / 2
                
                inside_table = False
                for t_box in table_bboxes:
                    if t_box[0] <= cx <= t_box[2] and t_box[1] <= cy <= t_box[3]:
                        inside_table = True
                        break
                
                # Nếu text thuộc về bảng, bỏ qua vì ta đã xử lý ở Bước 1
                if inside_table:
                    continue
                
                for line in block["lines"]:
                    bbox = line["bbox"]
                    
                    t_left = Inches(bbox[0] / 72)
                    t_top = Inches(bbox[1] / 72)
                    t_width = Inches(max((bbox[2] - bbox[0]), 20) / 72)
                    t_height = Inches(max((bbox[3] - bbox[1]), 10) / 72)
                    
                    txBox = slide.shapes.add_textbox(t_left, t_top, t_width, t_height)
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
                    p = tf.paragraphs[0]
                    
                    for span in line.get("spans", []):
                        run = p.add_run()
                        run.text = span.get("text", "")
                        run.font.size = PptxPt(max(span.get("size", 11), 8))
                        run.font.name = "Times New Roman"
                        
    doc.close()
    
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io.getvalue()


# ── Cấu hình trang Streamlit ────────────────────────────────────
st.set_page_config(
    page_title="Bộ chuyển đổi PDF đa năng",
    page_icon="🛠️",
    layout="centered"
)

# ── Thanh Menu bên cạnh (Sidebar) để chọn chức năng ──────────────
st.sidebar.title("⚙️ Bảng điều khiển")
conversion_type = st.sidebar.radio(
    "Chọn định dạng đầu ra:",
    ("Chuyển sang Word (.docx)", "Chuyển sang PowerPoint (.pptx)")
)

st.title("🛠️ Bộ chuyển đổi PDF đa năng")
st.write(f"Đang chọn chế độ: **{conversion_type}**")

# ── Upload file ──────────────────────────────────────────────────
uploaded_file = st.file_uploader("Chọn file PDF cần xử lý", type=["pdf"])
password = st.text_input("Mật khẩu PDF (nếu có, để trống nếu không)", type="password")

with st.expander("📐 Tùy chọn nâng cao (Giới hạn số trang)"):
    col1, col2 = st.columns(2)
    with col1:
        start_page = st.number_input("Trang bắt đầu (0 = từ đầu)", min_value=0, value=0)
    with col2:
        end_page = st.number_input("Trang kết thúc (0 = đến cuối)", min_value=0, value=0)

# ── Nút xử lý chính ──────────────────────────────────────────────
if uploaded_file is not None:
    st.info(f"✅ Đã nhận file: **{uploaded_file.name}**")

    if st.button("🚀 Bắt đầu chuyển đổi", use_container_width=True):
        with tempfile.TemporaryDirectory() as tmp_dir:
            pdf_path = os.path.join(tmp_dir, uploaded_file.name)

            with open(pdf_path, "wb") as f:
                f.write(uploaded_file.read())

            # Kiểm tra tính hợp lệ và đếm số trang của PDF
            try:
                doc = fitz.open(pdf_path)
                if doc.is_encrypted:
                    if not password or not doc.authenticate(password):
                        st.error("❌ PDF được bảo vệ — mật khẩu sai hoặc chưa nhập.")
                        st.stop()
                total_pages = len(doc)
                doc.close()
                st.write(f"📑 Tổng số trang PDF phát hiện: **{total_pages}**")
            except Exception as e:
                st.error(f"❌ Không đọc được file PDF: {e}")
                st.stop()

            # ── XỬ LÝ NHÁNH 1: CHUYỂN SANG WORD ────────────────────────
            if conversion_type == "Chuyển sang Word (.docx)":
                docx_name = os.path.splitext(uploaded_file.name)[0] + ".docx"
                docx_path = os.path.join(tmp_dir, docx_name)
                
                with st.spinner("⏳ Đang phân tích cấu trúc và chuyển sang Word..."):
                    try:
                        has_marker = pdf_has_page_marker(pdf_path, password if password else None)
                        start_0 = (start_page - 1) if start_page > 0 else 0
                        end_0 = end_page if end_page > 0 else None

                        cv = Converter(pdf_path, password=password if password else None)
                        cv.convert(docx_path, start=start_0, end=end_0)
                        cv.close()

                        post_process_word(docx_path, has_marker)

                        with open(docx_path, "rb") as f:
                            docx_bytes = f.read()

                        st.success("✅ Đã chuyển đổi sang Word thành công!")
                        st.download_button(
                            label="⬇️ Tải file Word về",
                            data=docx_bytes,
                            file_name=docx_name,
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )
                    except Exception as e:
                        st.error(f"❌ Lỗi khi chuyển đổi sang Word: {e}")

            # ── XỬ LÝ NHÁNH 2: CHUYỂN SANG POWERPOINT ──────────────────
            elif conversion_type == "Chuyển sang PowerPoint (.pptx)":
                pptx_name = os.path.splitext(uploaded_file.name)[0] + ".pptx"
                
                with st.spinner("⏳ Đang bóc tách chữ, dựng cấu trúc bảng sang PowerPoint..."):
                    try:
                        pptx_bytes = convert_pdf_
